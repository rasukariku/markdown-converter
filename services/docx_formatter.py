import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

_RE_PSEUDO_LIST = re.compile(r'^([a-zA-Z0-9]{1,4}[\.\)]|\([a-zA-Z0-9]{1,4}\)|[\-\*\+•●○■□▪▫►>])\s+')
_HR_PLACEHOLDER = '[[HR_PLACEHOLDER]]'

# Standard character count threshold for body text auto-justification
MIN_BODY_TEXT_LENGTH_FOR_JUSTIFY = 90


def format_docx_document(doc: Document) -> Document:
    """
    Apply comprehensive formatting to a DOCX document including styles,
    headings, code blocks, lists, and dynamic alignment.
    """
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    try:
        hlink_style = doc.styles['Hyperlink']
        hlink_style.font.name = 'Times New Roman'
        hlink_style.font.color.rgb = RGBColor(0, 0, 0)
        if hlink_style.element.rPr is not None:
            color_el = hlink_style.element.rPr.find(qn('w:color'))
            if color_el is not None and qn('w:themeColor') in color_el.attrib:
                del color_el.attrib[qn('w:themeColor')]
            rFonts = hlink_style.element.rPr.find(qn('w:rFonts'))
            if rFonts is not None:
                for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme']:
                    if qn(attr) in rFonts.attrib:
                        del rFonts.attrib[qn(attr)]
    except KeyError:
        pass
    
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = Cm(2.54), Cm(2.54), Cm(2.54), Cm(2.54)
    
    settings = doc.settings.element
    math_pr = settings.find(qn('m:mathPr'))
    if math_pr is None:
        math_pr = OxmlElement('m:mathPr')
        settings.append(math_pr)
        
    def_jc = math_pr.find(qn('m:defJc'))
    if def_jc is None:
        def_jc = OxmlElement('m:defJc')
        math_pr.append(def_jc)
    def_jc.set(qn('m:val'), 'left')
    
    compat = settings.find(qn('w:compat'))
    if compat is None:
        compat = OxmlElement('w:compat')
        settings.append(compat)
        
    compat_setting = OxmlElement('w:compatSetting')
    compat_setting.set(qn('w:name'), 'compatibilityMode')
    compat_setting.set(qn('w:uri'), 'http://schemas.microsoft.com/office/word')
    compat_setting.set(qn('w:val'), '15')
    compat.append(compat_setting)
    
    _process_paragraphs(doc)
    _process_tables(doc)
    
    return doc


def _process_paragraphs(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    removal_queue = []
    
    for i, para in enumerate(paragraphs):
        text_clean = para.text.strip()
        style_name = para.style.name
        
        has_math = bool(para._element.findall('.//' + qn('m:oMath'))) or bool(para._element.findall('.//' + qn('m:oMathPara')))
        has_drawing = bool(para._element.findall('.//' + qn('w:drawing')))
        has_soft_break = bool(para._element.findall('.//' + qn('w:br')))
        
        is_heading = style_name.startswith('Heading')
        is_list = ('List' in style_name or 'Bullet' in style_name or 'Compact' in style_name or bool(para._element.findall('.//' + qn('w:numPr'))))
        is_code = 'Source Code' in style_name or 'Code' in style_name or 'Preformatted' in style_name
        is_pseudo_list = bool(_RE_PSEUDO_LIST.match(text_clean))
        is_quote = 'Quote' in style_name or 'Block Text' in style_name
        is_hr = _HR_PLACEHOLDER in text_clean
        
        # 1. Process Horizontal Dividers
        if is_hr:
            p_el = para._element
            for run in list(para.runs):
                p_el.remove(run._element)
            pPr = p_el.get_or_add_pPr()
            old_bdr = pPr.find(qn('w:pBdr'))
            if old_bdr is not None:
                pPr.remove(old_bdr)
            pBdr = OxmlElement('w:pBdr')
            top_border = OxmlElement('w:top')
            top_border.set(qn('w:val'), 'single')
            top_border.set(qn('w:sz'), '12')
            top_border.set(qn('w:space'), '0')
            top_border.set(qn('w:color'), '000000')
            pBdr.append(top_border)
            pPr.append(pBdr)
            try:
                para.style = doc.styles['Normal']
            except KeyError:
                pass
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            para.paragraph_format.line_spacing = Pt(1)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.left_indent = Pt(0)
            para.paragraph_format.first_line_indent = Pt(0)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            has_border = True
            text_clean = ""
        else:
            has_border = False
            
        if not text_clean and not has_math and not has_drawing and not is_code and not has_border:
            removal_queue.append(para)
            continue

        # 2. Unified Paragraph Layout, Spacing, and Alignment Decision Tree
        if is_heading:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.line_spacing = 1.2
            
            heading_size = Pt(12)
            space_before = Pt(10)
            space_after = Pt(4)
            is_italic_heading = False
            
            if 'Heading 1' in style_name:
                heading_size = Pt(18)
                space_before = Pt(16)
                space_after = Pt(6)
            elif 'Heading 2' in style_name:
                heading_size = Pt(15)
                space_before = Pt(13)
                space_after = Pt(6)
            elif 'Heading 3' in style_name:
                heading_size = Pt(13)
                space_before = Pt(10)
                space_after = Pt(4)
            elif 'Heading 4' in style_name:
                heading_size = Pt(12)
                space_before = Pt(8)
                space_after = Pt(4)
            elif 'Heading 5' in style_name:
                heading_size = Pt(12)
                space_before = Pt(6)
                space_after = Pt(2)
                is_italic_heading = True
            elif 'Heading 6' in style_name:
                heading_size = Pt(11)
                space_before = Pt(6)
                space_after = Pt(2)
                is_italic_heading = True

            para.paragraph_format.space_before = space_before
            para.paragraph_format.space_after = space_after

        elif is_code:
            # Coding format: Monospace, single line spacing, indented block
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(8)
            para.paragraph_format.left_indent = Pt(18)
            para.paragraph_format.first_line_indent = Pt(0)

        elif is_list or is_pseudo_list:
            # Bullet/Numbered lists: Strictly LEFT-aligned
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_before = Pt(0)
            
            next_is_list = False
            if i + 1 < len(paragraphs):
                next_para = paragraphs[i + 1]
                next_is_list = ('List' in next_para.style.name or 'Bullet' in next_para.style.name or 'Compact' in next_para.style.name or bool(next_para._element.findall('.//' + qn('w:numPr'))))
            para.paragraph_format.space_after = Pt(4) if next_is_list else Pt(8)
            
            ilvl_nodes = para._element.findall('.//' + qn('w:ilvl'))
            level = int(ilvl_nodes[0].get(qn('w:val'))) if ilvl_nodes else 0
            para.paragraph_format.left_indent = Pt(36 + (level * 36))
            para.paragraph_format.first_line_indent = Pt(-18)

        elif is_quote:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(8)
            para.paragraph_format.left_indent = Cm(1.5)
            para.paragraph_format.right_indent = Cm(1.5)

        else:
            # Regular text paragraphs
            if not has_border:
                para.style = doc.styles['Normal']
                para.paragraph_format.line_spacing = 1.5
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(8)
                para.paragraph_format.left_indent = Pt(0)
                para.paragraph_format.first_line_indent = Pt(0)

                # Dynamic Justify: Apply strictly to multi-line narratives without soft breaks
                word_count = len(text_clean.split())
                is_long_narrative = (len(text_clean) >= MIN_BODY_TEXT_LENGTH_FOR_JUSTIFY) and (word_count >= 12)
                
                if is_long_narrative and not has_soft_break:
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                else:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # 3. Unified Typography & Run Formatting
        for run in para.runs:
            if run._element.findall('.//' + qn('m:oMath')):
                continue
            
            rPr = run._element.get_or_add_rPr()
            
            # Detect inline code formatting produced by Pandoc verbatim markup
            is_inline_code = (
                'Verbatim' in run.style.name or 
                'Code' in run.style.name or 
                'Source' in run.style.name or 
                run.font.name in ['Consolas', 'Courier New']
            )

            if is_code or is_inline_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(10.0) if is_code else Pt(10.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is None:
                    rFonts = OxmlElement('w:rFonts')
                    rPr.append(rFonts)
                rFonts.set(qn('w:ascii'), 'Consolas')
                rFonts.set(qn('w:hAnsi'), 'Consolas')
                rFonts.set(qn('w:cs'), 'Consolas')
            elif is_heading:
                run.font.name = 'Times New Roman'
                run.font.size = heading_size
                run.font.bold = True
                if is_italic_heading:
                    run.font.italic = True
                run.font.color.rgb = RGBColor(0, 0, 0)
            else:
                run.font.name = 'Times New Roman'
                if run.font.size is None:
                    run.font.size = Pt(12)
                if run.font.color.rgb is None and 'Hyperlink' not in run.style.name:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                
                color_el = rPr.find(qn('w:color'))
                if color_el is not None and qn('w:themeColor') in color_el.attrib:
                    del color_el.attrib[qn('w:themeColor')]
                if 'Hyperlink' in run.style.name or 'Hyperlink' in style_name:
                    run.font.underline = True

            lang_el = rPr.find(qn('w:lang'))
            if lang_el is None:
                lang_el = OxmlElement('w:lang')
                rPr.append(lang_el)
            lang_el.set(qn('w:val'), 'id-ID')


def _process_tables(doc: Document) -> None:
    for table in doc.tables:
        try:
            table.style = 'Table Grid'
        except KeyError:
            try:
                table.style = 'TableGrid'
            except KeyError:
                pass
                
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    orig_align = para.alignment
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.5
                    para.alignment = orig_align if orig_align is not None else WD_ALIGN_PARAGRAPH.LEFT
                    for run in para.runs:
                        if run._element.findall('.//' + qn('m:oMath')):
                            continue
                        run.font.name = 'Times New Roman'
                        if run.font.size is None:
                            run.font.size = Pt(12)
                        if run.font.color.rgb is None and 'Hyperlink' not in run.style.name:
                            run.font.color.rgb = RGBColor(0, 0, 0)
                            
        tbl_element = table._element
        next_element = tbl_element.getnext()
        if next_element is not None and next_element.tag == qn('w:p'):
            next_para = Paragraph(next_element, doc)
            next_para.paragraph_format.space_before = Pt(8)